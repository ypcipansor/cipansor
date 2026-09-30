#!/usr/bin/env python3
"""Uji regresi untuk perkakas dokumen (Devin review: "Document tooling lacks regression coverage").

Kenapa berkas ini ada: `build_docs.py` dan `check_docs.py` adalah gerbang yang memutuskan
apakah naskah boleh diserahkan, dan keduanya menutup kesalahan yang tidak terlihat dari kode
keluar — tabel yang terbit kosong, halaman nyaris kosong, angka yang tidak cocok dengan
`facts.json`, rute yang tidak terdaftar. Sebelum ini tak satu pun dari jalur itu punya uji,
jadi perubahan pada perkakas bisa melumpuhkan gerbang tanpa ada yang menangkapnya.

Uji di sini memakai berkas contoh kecil dan `monkeypatch` untuk Prisma/router palsu, sehingga
tidak perlu repo penuh maupun LibreOffice. Jalankan:

    python -m unittest discover -s .claude/skills/dokumen-aplikasi-cipansor/scripts -p 'test_*.py'

atau langsung:

    python .claude/skills/dokumen-aplikasi-cipansor/scripts/test_tooling.py
"""
from __future__ import annotations

import importlib
import sys
import unittest
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

check_docs = importlib.import_module("check_docs")
build_docs = importlib.import_module("build_docs")
collect_facts = importlib.import_module("collect_facts")


class NormRouteTest(unittest.TestCase):
    def test_parameters_collapse_to_one_shape(self):
        self.assertEqual(check_docs.norm_route("/students/:id"), "/api/students/:p")
        self.assertEqual(check_docs.norm_route("/students/{id}"), "/api/students/:p")
        self.assertEqual(check_docs.norm_route("/students/[id]"), "/api/students/:p")

    def test_trailing_slash_and_query_are_stripped(self):
        self.assertEqual(check_docs.norm_route("/students/"), "/api/students")
        self.assertEqual(check_docs.norm_route("/students?page=1"), "/api/students")

    def test_punctuation_after_a_path_is_dropped(self):
        # A route cited at the end of a sentence must not fail on the period.
        self.assertEqual(check_docs.norm_route("/students/:id."), "/api/students/:p")

    def test_api_prefix_is_not_doubled(self):
        self.assertEqual(check_docs.norm_route("/api/v1/students"), "/api/v1/students")


class SplitRowTest(unittest.TestCase):
    def test_cells_split_on_pipes(self):
        self.assertEqual(check_docs.split_row("| a | b | c |"), ["a", "b", "c"])

    def test_a_pipe_inside_backticks_does_not_split(self):
        self.assertEqual(check_docs.split_row("| `a|b` | c |"), ["`a|b`", "c"])


class CountMdTablesTest(unittest.TestCase):
    def test_counts_each_separator_row_once(self):
        md = "| a | b |\n|---|---|\n| 1 | 2 |\n\n| c |\n|---|\n| 3 |\n"
        self.assertEqual(build_docs.count_md_tables(md), 2)

    def test_ignores_tables_inside_code_fences(self):
        md = "```\n| a | b |\n|---|---|\n```\n\n| c |\n|---|\n| 3 |\n"
        self.assertEqual(build_docs.count_md_tables(md), 1)

    def test_zero_when_there_are_no_tables(self):
        self.assertEqual(build_docs.count_md_tables("just prose\n"), 0)


class CheckRoutesTest(unittest.TestCase):
    """The route gate is what keeps a document from citing an API that does not exist."""

    def _run(self, line: str, routes: list[tuple[str, str]]):
        r = check_docs.Report()
        idx = [{"method": m, "path": p, "module": "x"} for m, p in routes]
        orig = collect_facts.route_index
        collect_facts.route_index = lambda repo: idx  # type: ignore[assignment]
        try:
            check_docs.check_routes(line, [line], Path("."), r)
        finally:
            collect_facts.route_index = orig  # type: ignore[assignment]
        return r

    def test_a_registered_route_passes(self):
        r = self._run("GET /api/students", [("GET", "/api/students")])
        self.assertEqual([i for i in r.items if i[0] == "ERROR"], [])

    def test_an_unregistered_route_is_an_error(self):
        r = self._run("GET /api/students", [("GET", "/api/teachers")])
        errors = [i for i in r.items if i[0] == "ERROR"]
        self.assertEqual(len(errors), 1)
        self.assertEqual(errors[0][2], "rute-tak-ada")

    def test_a_wrong_method_is_an_error_and_says_so(self):
        r = self._run("DELETE /api/students", [("GET", "/api/students")])
        errors = [i for i in r.items if i[0] == "ERROR"]
        self.assertEqual(len(errors), 1)
        self.assertIn("metode lain", errors[0][3])

    def test_parameters_match_across_notations(self):
        r = self._run("GET /api/students/{id}", [("GET", "/api/students/:id")])
        self.assertEqual([i for i in r.items if i[0] == "ERROR"], [])


class CheckNumbersTest(unittest.TestCase):
    """The number gate compares prose against facts.json, but a revision-history
    row records the number as it was at that commit — it must not be rewritten."""

    FACTS = {
        "api": {"module_count": 94, "handler_count_approx": 1412},
        "prisma": {"model_count": 289, "enum_count": 157, "role_code_count": 53},
        "web": {"page_count": 435},
        "jobs": {"cron_entries": 15, "scheduled_job_files": ["a.job.ts"], "count": 1,
                 "job_files_not_scheduled": []},
        "env": {"total_keys": 40},
    }

    def _run(self, line: str):
        r = check_docs.Report()
        check_docs.check_numbers([line], self.FACTS, r)
        return [i for i in r.items if i[0] == "ERROR"]

    def test_a_stale_current_number_is_an_error(self):
        errors = self._run("Aplikasi punya 93 modul API.")
        self.assertEqual(len(errors), 1)
        self.assertEqual(errors[0][2], "angka-salah")

    def test_a_correct_number_passes(self):
        self.assertEqual(self._run("Aplikasi punya 94 modul API."), [])

    def test_a_revision_history_row_is_exempt(self):
        # History: "~1.413 hulu rute API (dari ~1.412)" was true at that commit.
        self.assertEqual(self._run("| 0.7 | 30 September 2026 | ~1.413 hulu rute API | Agen |"), [])


class VerifyDocxTest(unittest.TestCase):
    """A .docx whose structure drifted from the Markdown must be reported, not passed."""

    @staticmethod
    def _docx(tmp: Path, *, tables: int, with_heading_in_cell: bool = False, toc: bool = True):
        from docx import Document

        d = Document()
        d.add_table(rows=1, cols=2)  # [0] is the cover metadata table, skipped
        for _ in range(tables):
            t = d.add_table(rows=2, cols=2)
            t.cell(0, 0).text = "Kepala"
            t.cell(0, 1).text = "Nilai"
            t.cell(1, 0).text = "isi"
            t.cell(1, 1).text = "isi"
            if with_heading_in_cell:
                t.cell(1, 0).paragraphs[0].style = d.styles["Heading 1"]
        if toc:
            d.add_paragraph("Daftar Isi")
        out = tmp / "doc.docx"
        d.save(str(out))
        return out

    def test_matching_structure_has_no_problems(self):
        import tempfile

        with tempfile.TemporaryDirectory() as td:
            p = self._docx(Path(td), tables=1)
            self.assertEqual(build_docs.verify_docx(p, "| a | b |\n|---|---|\n| 1 | 2 |\n", 0), [])

    def test_missing_table_is_reported(self):
        import tempfile

        with tempfile.TemporaryDirectory() as td:
            p = self._docx(Path(td), tables=1)
            problems = build_docs.verify_docx(p, "| a | b |\n|---|---|\n| 1 | 2 |\n\n| c | d |\n|---|---|\n| 3 | 4 |\n", 0)
            self.assertTrue(any("jumlah tabel" in x for x in problems), problems)

    def test_heading_inside_a_cell_is_reported(self):
        import tempfile

        with tempfile.TemporaryDirectory() as td:
            p = self._docx(Path(td), tables=1, with_heading_in_cell=True)
            problems = build_docs.verify_docx(p, "| a | b |\n|---|---|\n| 1 | 2 |\n", 0)
            self.assertTrue(any("judul bab" in x for x in problems), problems)

    def test_missing_daftar_isi_is_reported(self):
        import tempfile

        with tempfile.TemporaryDirectory() as td:
            p = self._docx(Path(td), tables=1, toc=False)
            problems = build_docs.verify_docx(p, "| a | b |\n|---|---|\n| 1 | 2 |\n", 0)
            self.assertTrue(any("Daftar Isi" in x for x in problems), problems)


class VerifyPdfTest(unittest.TestCase):
    """The PDF gate must notice a near-empty page and a table whose header did not render."""

    @staticmethod
    def _pdf(tmp: Path, pages: list[str]):
        import pymupdf

        doc = pymupdf.open()
        for text in pages:
            page = doc.new_page()
            page.insert_text((72, 72), text)
        out = tmp / "doc.pdf"
        doc.save(str(out))
        doc.close()
        return out

    def test_a_table_header_missing_from_the_pdf_is_reported(self):
        import tempfile

        with tempfile.TemporaryDirectory() as td:
            p = self._pdf(Path(td), ["Halaman dengan teks yang cukup panjang untuk diperiksa mesin dokumen ini."])
            md = "| Modul | Jumlah |\n|---|---|\n| api | 3 |\n"
            problems = build_docs.verify_pdf(p, md)
            self.assertTrue(any("Modul" in x for x in problems), problems)

    def test_a_present_header_passes(self):
        import tempfile

        with tempfile.TemporaryDirectory() as td:
            p = self._pdf(Path(td), ["Modul Jumlah api 3 — teks yang cukup panjang untuk diperiksa mesin."])
            md = "| Modul | Jumlah |\n|---|---|\n| api | 3 |\n"
            self.assertEqual(build_docs.verify_pdf(p, md), [])


if __name__ == "__main__":
    unittest.main(verbosity=2)
