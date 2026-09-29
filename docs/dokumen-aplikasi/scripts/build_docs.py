#!/usr/bin/env python3
"""Markdown -> .docx (sampul, daftar isi, tabel, diagram Mermaid) dan/atau .md.

Direkonstruksi sesuai antarmuka yang didokumentasikan SKILL.md (berkas asli
hilang saat sesi). Memakai pandoc, python-docx, dan mmdc (mermaid-cli) bila ada.

Pemakaian:
    python build_docs.py naskah.md --out keluaran --format docx \
      --title "Dokumen Teknis Aplikasi" --subtitle "Sistem Informasi Cipansor" \
      --version 0.1 --status Draf --commit <hash> --logo logo.svg
"""
from __future__ import annotations

import argparse
import datetime as dt
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent


def tanggal_id(d: dt.date) -> str:
    bulan = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli",
             "Agustus", "September", "Oktober", "November", "Desember"]
    return f"{d.day} {bulan[d.month - 1]} {d.year}"


# --------------------------------------------------------------------------- mermaid

def render_mermaid(md: str, out: Path) -> tuple[str, int, int]:
    """Ganti blok ```mermaid dengan gambar PNG. Kembalikan (md2, ok, gagal)."""
    pattern = re.compile(r"```mermaid\s*\n(.*?)```", re.S)
    mmdc = shutil.which("mmdc") or str(Path(os.environ.get("WORKDIR", ".")) / "node_modules/.bin/mmdc")
    ok = bad = 0
    diagrams_dir = out / "diagrams"
    diagrams_dir.mkdir(parents=True, exist_ok=True)

    def repl(m: re.Match) -> str:
        nonlocal ok, bad
        code = m.group(1).strip()
        idx = ok + bad + 1
        mmd = diagrams_dir / f"diagram-{idx}.mmd"
        png = diagrams_dir / f"diagram-{idx}.png"
        mmd.write_text(code, encoding="utf-8")
        pconf = out / "puppeteer.json"
        if not pconf.exists():
            pconf.write_text('{"args":["--no-sandbox","--disable-setuid-sandbox","--disable-dev-shm-usage"]}', encoding="utf-8")
        r = subprocess.run(
            [mmdc, "-i", str(mmd), "-o", str(png), "-b", "white", "-p", str(pconf), "--size", "1200"],
            capture_output=True, text=True,
            env={**os.environ},
        )
        if r.returncode == 0 and png.exists():
            ok += 1
            return f"![diagram]({png.as_posix()}){{width=15cm}}"
        bad += 1
        print(f"    diagram {idx} GAGAL:\n{r.stderr[-500:]}", file=sys.stderr)
        return m.group(0)

    return pattern.sub(repl, md), ok, bad


# --------------------------------------------------------------------------- reference docx

def make_reference(path: Path, accent: str) -> None:
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm

    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2.2)
        s.left_margin = s.right_margin = Cm(2.2)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for name, size in (("Title", 26), ("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
        try:
            st = doc.styles[name]
            st.font.size = Pt(size)
            st.font.color.rgb = RGBColor.from_string(accent if name != "Title" else "1A1A1A")
            st.font.bold = True
        except KeyError:
            pass
    for p in doc.paragraphs:
        p.paragraph_format.space_after = Pt(6)
    doc.save(str(path))


# --------------------------------------------------------------------------- postprocess

def postprocess(docx_path: Path, a) -> None:
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import qn, nsdecls

    doc = Document(str(docx_path))
    accent = a.accent

    front: list = []  # elemen sampul + daftar isi, dipindah ke awal setelah dibuat

    def front_add(el):
        front.append(el)

    def _field(par, code, size=8.5, color="777777"):
        r = par.add_run()
        fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "begin")
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {code} "
        fc2 = OxmlElement("w:fldChar"); fc2.set(qn("w:fldCharType"), "end")
        r._r.append(fc); r._r.append(it); r._r.append(fc2)
        r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string(color)

    # ---- sampul
    cover = doc.add_paragraph()
    front_add(cover._p)
    title = cover.add_run(a.title)
    title.font.size = Pt(28); title.font.bold = True
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.space_before = Pt(120)
    if a.subtitle:
        sp = doc.add_paragraph(); front_add(sp._p)
        r = sp.add_run(a.subtitle); r.font.size = Pt(16); r.font.color.rgb = RGBColor.from_string("555555")
        sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = [("Versi", a.version), ("Tanggal", a.date_text), ("Status", a.status)]
    if a.commit:
        meta.append(("Basis kode", f"commit {a.commit}"))
    meta.append(("Klasifikasi", a.classification))
    mt = doc.add_table(rows=len(meta), cols=2)
    mt.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate(meta):
        c0, c1 = mt.rows[i].cells
        rr = c0.paragraphs[0].add_run(k); rr.font.bold = True; rr.font.size = Pt(11)
        rr = c1.paragraphs[0].add_run(v); rr.font.size = Pt(11)
    front_add(mt._tbl)
    pb = doc.add_paragraph(); pb.add_run().add_break(WD_BREAK.PAGE); front_add(pb._p)

    # ---- TOC field
    h = doc.add_paragraph(); front_add(h._p)
    hr = h.add_run("Daftar Isi"); hr.font.size = Pt(18); hr.font.bold = True
    hr.font.color.rgb = RGBColor.from_string(accent)
    tocp = doc.add_paragraph(); front_add(tocp._p)
    for typ, txt in (("begin", None), ("instr", ' TOC \\o "1-3" \\h \\z \\u '), ("separate", None),
                     ("text", "Daftar isi dibuat otomatis — klik kanan > Perbarui bidang."), ("end", None)):
        r = tocp.add_run()
        if typ == "instr":
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; r._r.append(it)
        elif typ == "text":
            r.text = txt
        else:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), typ)
            if typ == "begin":
                fc.set(qn("w:dirty"), "true")
            r._r.append(fc)
    pb2 = doc.add_paragraph(); pb2.add_run().add_break(WD_BREAK.PAGE); front_add(pb2._p)

    # ---- tabel bergaris
    for tbl in doc.tables:
        tblPr = tbl._tbl.tblPr
        for old in tblPr.findall(qn("w:tblBorders")):
            tblPr.remove(old)
        tblPr.append(parse_xml(
            f'<w:tblBorders {nsdecls("w")}>' + "".join(
                f'<w:{s} w:val="single" w:sz="4" w:space="0" w:color="C9CED3"/>'
                for s in ("top", "left", "bottom", "right", "insideH", "insideV")) + "</w:tblBorders>"))
        for ri, row in enumerate(tbl.rows):
            trPr = row._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:cantSplit"))
            if ri == 0:
                trPr.append(OxmlElement("w:tblHeader"))
            for cell in row.cells:
                if ri == 0:
                    cell._tc.get_or_add_tcPr().append(
                        parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{accent}"/>'))
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(9.5)
                        if ri == 0:
                            r.font.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # ---- header/footer
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    hp = sec.header.paragraphs[0]; hp.text = ""
    r = hp.add_run(f"{a.title} — {a.subtitle}" if a.subtitle else a.title)
    r.font.size = Pt(8.5); r.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp = sec.footer.paragraphs[0]; fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = fp.add_run(f"{a.org} · v{a.version} · Halaman ")
    rr.font.size = Pt(8.5); rr.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
    _field(fp, "PAGE")
    rr = fp.add_run(" dari "); rr.font.size = Pt(8.5); rr.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
    _field(fp, "NUMPAGES")

    # gambar maksimal 15 cm
    for shape in doc.inline_shapes:
        if shape.width > Cm(15):
            ratio = Cm(15) / shape.width
            shape.width = Cm(15)
            shape.height = int(shape.height * ratio)

    # pindahkan sampul + daftar isi ke awal badan dokumen (sebelum Riwayat Revisi)
    body = doc.element.body
    first = None
    for el in list(body):
        if el.tag != qn("w:sectPr"):
            first = el
            break
    for el in front:
        if first is not None:
            first.addprevious(el)
        else:
            body.append(el)

    doc.core_properties.title = f"{a.title} — {a.subtitle}" if a.subtitle else a.title
    doc.core_properties.author = a.author or a.org
    doc.core_properties.language = "id-ID"
    doc.save(str(docx_path))


def set_update_fields(docx_path: Path) -> None:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document(str(docx_path))
    el = OxmlElement("w:updateFields"); el.set(qn("w:val"), "true")
    doc.settings.element.append(el)
    doc.save(str(docx_path))


# --------------------------------------------------------------------------- main

def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("source")
    ap.add_argument("--out", default="out")
    ap.add_argument("--format", choices=["docx", "md", "both"], default="docx")
    ap.add_argument("--name")
    ap.add_argument("--title", required=True)
    ap.add_argument("--subtitle", default="")
    ap.add_argument("--version", default="1.0")
    ap.add_argument("--status", default="Draf")
    ap.add_argument("--author", default="")
    ap.add_argument("--commit", default="")
    ap.add_argument("--org", default="Yayasan Pesantren Cipansor")
    ap.add_argument("--classification", default="Internal — Yayasan Pesantren Cipansor")
    ap.add_argument("--logo", default="")
    ap.add_argument("--accent", default="0B5D3B")
    ap.add_argument("--resource-path", default="")
    ap.add_argument("--no-toc-update", action="store_true")
    a = ap.parse_args()
    a.date_text = tanggal_id(dt.date.today())

    src = Path(a.source).resolve()
    if not src.is_file():
        print(f"ERROR: {src} tidak ada", file=sys.stderr); return 2
    out = Path(a.out).resolve(); out.mkdir(parents=True, exist_ok=True)
    name = a.name or src.stem
    md = src.read_text(encoding="utf-8")

    left = re.findall(r"\[ISI[^\]]*\]|\bTODO\b|\bTBD\b|\bLOREM\b", md)
    if left:
        print(f"PERINGATAN: {len(left)} penanda isian tersisa, mis. {left[0]!r}. Dokumen belum final.", file=sys.stderr)

    if a.format in ("md", "both"):
        shutil.copyfile(src, out / f"{name}.md")
        print(f"OK  {out / (name + '.md')}")

    if a.format in ("docx", "both"):
        pandoc = shutil.which("pandoc") or os.environ.get("PANDOC")
        if not pandoc:
            print("ERROR: pandoc tidak ditemukan", file=sys.stderr); return 2
        md2, ok, bad = render_mermaid(md, out)
        if ok or bad:
            print(f"    diagram: {ok} dirender, {bad} gagal")
        with tempfile.TemporaryDirectory() as td:
            tdp = Path(td); ref = tdp / "ref.docx"; make_reference(ref, a.accent)
            tmp_md = tdp / "in.md"; tmp_md.write_text(md2, encoding="utf-8")
            target = out / f"{name}.docx"
            rp = os.pathsep.join(p for p in (a.resource_path, str(src.parent), str(out)) if p)
            cmd = [pandoc, str(tmp_md), "-f", "commonmark_x-smart", "-o", str(target),
                   f"--reference-doc={ref}", f"--resource-path={rp}"]
            r = subprocess.run(cmd, capture_output=True, text=True)
            if r.returncode != 0:
                print(f"ERROR pandoc:\n{r.stderr}", file=sys.stderr); return 1
            if r.stderr.strip():
                print("    pandoc:", r.stderr.strip()[:600])
        postprocess(target, a)
        set_update_fields(target)
        print("    daftar isi: bidang ditandai untuk diperbarui; Word akan mengisinya saat dibuka")
        print(f"OK  {target}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
